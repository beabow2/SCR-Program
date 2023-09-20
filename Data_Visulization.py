#load package
import pandas as pd
import os
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
class Visualization:
    #initialize
    def __init__(self):
        pass

    #data loading
    def data_loading(self,filepath):
        data = pd.read_csv(filepath,index_col=0)
        return(data)

    def line_chart(self,data,filename,savepath):
        new_folder = "chart"
        new_savepath = os.path.join(savepath, new_folder)
        if not os.path.exists(new_savepath):
            os.makedirs(new_savepath)
        #draw CS+ and CS-
        fig1, ax1 = plt.subplots()
        ax1.set_ylim(0,1)
        ax1.plot(data['CSP'], label='CS+')
        ax1.plot(data['CSM'], label='CS-')
        ax1.set_title('CS+ and CS-')
        ax1.legend()
        image_path = os.path.join(new_savepath, f'CS+ and CS-_{filename}.png')
        fig1.savefig(image_path)
        plt.close(fig1)
        #draw diff
        fig2, ax2 = plt.subplots()
        ax2.set_ylim(-0.5,0.5)
        ax2.plot(data['difference'], label='difference')
        ax2.set_title('Difference')
        image_path = os.path.join(new_savepath, f'difference_{filename}.png')
        fig2.savefig(image_path)
        plt.close(fig2)

    def add_chart(self, data, sheet_name, filename, savepath):
        new_folder = "excel_data"
        new_savepath = os.path.join(savepath, new_folder)
        if not os.path.exists(new_savepath):
            os.makedirs(new_savepath)
        excel_path = os.path.join(new_savepath, f'{filename}.xlsx')

        # Save the Excel file first
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            data.to_excel(writer, sheet_name=sheet_name, index=True)
            writer.save()  # Explicitly save the Excel file

        # Now, read the Excel file and add the images
        book = load_workbook(excel_path)
        sheet = book[sheet_name]  # Get the sheet

        # Add CS+ and CS- image
        image_path = os.path.join(savepath, "chart", f'CS+ and CS-_{filename}.png')
        CS_image = Image(image_path)
        CS_image.width = CS_image.width
        CS_image.height = CS_image.height
        sheet.add_image(CS_image, 'G5')

        # Add difference image
        image_path = os.path.join(savepath, "chart", f'difference_{filename}.png')
        diff_image = Image(image_path)
        diff_image.width = diff_image.width
        diff_image.height = diff_image.height
        sheet.add_image(diff_image, 'S5')

        # Save the changes
        book.save(excel_path)
    def batch_visualization(self,folderpath):

        for filename in os.listdir(folderpath):
            print(filename)
            # Data loading
            file_path = os.path.join(folderpath, filename)
            if not file_path.endswith('.csv'):
                continue
            data = self.data_loading(file_path)
            self.line_chart(data,filename,folderpath)
            self.add_chart(data,filename[:15],filename,folderpath)



if __name__ == "__main__":
    
    app = Visualization()
    #day1
    app.batch_visualization(r"D:\data\Day1\leda\result")
    #day2
    app.batch_visualization(r"D:\data\Day2\result")










